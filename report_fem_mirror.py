#!/usr/bin/env python3
"""
report_fem_mirror.py
====================
Generación de informes Word (.docx) del análisis FEM de superficie de espejo.

Uso desde la GUI (gui_fem_mirror.py):
    data = build_report_data(...)          # tras el análisis
    generate_report(data, 'informe.docx', lang='es', units='mm', logo_png=bytes)

Requiere python-docx:   pip install python-docx
"""

import datetime as _dt
import io
import math
import os

import numpy as np

import fem_mirror_analysis as fma


# ═══════════════════════════════════════════════════════════════════════════════
#  TEXTOS (ES / EN)
# ═══════════════════════════════════════════════════════════════════════════════

TXT = {
    'es': {
        'title': "Informe de análisis FEM de superficie de espejo",
        'header': "Análisis FEM – superficie de espejo",
        'footer': "Generado con fem-mirror-analysis",
        'page': "Página",
        's_setup': "1. Datos del análisis",
        's_rbm': "2. Superficie tras eliminar el movimiento de sólido rígido",
        's_fit': "3. Residuo tras eliminar {n} términos de {fam}",
        's_notes': "{k}. Observaciones",
        's_opt': "{k}. Calidad óptica",
        's_iso': "{k}. Tolerancias de forma ISO 10110-5",
        'combo': "Combinación de casos",
        'opt_text': ("Error de frente de onda en reflexión WFE = 2·cos θ·(error de superficie), "
                     "con λ = {lam:g} nm e incidencia θ = {inc:g}°, sin pistón ni tilt. Strehl "
                     "por la aproximación de Maréchal S ≈ exp[−(2πσ_W/λ)²] y numérico a partir "
                     "de la PSF calculada por FFT de la pupila."),
        'opt_surf': "Superficie sin sólido rígido", 'opt_res': "Residuo del ajuste",
        'opt_h': ["", "σ_W [nm]", "σ_W [λ]", "Strehl (Maréchal)"],
        'opt_fft': "Strehl numérico (PSF por FFT, {on})",
        'on_res': "residuo", 'on_surf': "superficie",
        'opt_note_lo': ("Strehl < 0,3: la aproximación de Maréchal no es fiable; vale el valor "
                        "numérico."),
        'opt_note_dl': ("Superficie sin sólido rígido con Strehl ≥ 0,8 (σ_W ≤ λ/14): limitada por "
                        "difracción a esta longitud de onda (criterio de Maréchal)."),
        'opt_note_ndl': ("Superficie sin sólido rígido con Strehl < 0,8: no es limitada por "
                         "difracción a esta λ sin corregir los términos ajustados."),
        'iso_text': ("Descomposición de la desviación de forma según ISO 10110-5:2015 (Tabla B.2) "
                     "sobre la apertura, tras quitar tilt. Valores de superficie en nm y en "
                     "fringe spacings (1 fringe = λ/2 = {fr:g} nm)."),
        'iso_h': ["Magnitud", "Función", "nm", "fringes"],
        'iso_rows': [("A · potencia", "PV(f_WS)"), ("B · irregularidad", "PV(f_WI)"),
                     ("C · irregularidad rot. invariante", "PV(f_WRI)"),
                     ("PV total", "PV(f_WD)"), ("RMSt · total", "rms(f_WD)"),
                     ("RMSi · irregularidad", "rms(f_WI)"),
                     ("RMSa · irreg. asimétrica", "rms(f_WRV)"),
                     ("PVr · PV robusto", "PVr(f_WD)"),
                     ("PV99.8% · PV recortado", "PV%(f_WD)"),
                     ("PV99.8% · irregularidad recortada", "PV%(f_WI)")],
        'iso_slope': "Pendiente de f_WI (muestreo de la malla FEM): RMS {r:.3g} µrad · máx {m:.3g} µrad.",
        'iso_notation': "Valores medidos en notación ISO: {s}",
        'iso_check': "Comprobación frente a la tolerancia {s}",
        'iso_ch_h': ["Magnitud", "Medido [nm]", "Tolerancia [nm]", "Resultado"],
        'ok': "Cumple", 'nok': "No cumple",
        'iso_method': ("Método: funciones de ISO 14999-4:2007 (3.2–3.3) con RMS ponderado por "
                       "área nodal. Tilt = plano best-fit (incluye pistón); esfera best-fit "
                       "aproximada paraxialmente por el término ρ²; parte rotacionalmente "
                       "invariante (cuyo óptimo rms es la media azimutal) aproximada por "
                       "polinomios pares en ρ hasta ρ^{nr}; PVr según Evans (Opt. Eng. 48(4), "
                       "2009), PV del ajuste de 37 Zernike Fringe + 3·RMS del residuo, acotado "
                       "entre 6·RMS del residuo y el PV (Murphy, Proc. SPIE 12778, 2023); PV99.8% "
                       "= PV del 99,8 % del área quitando de cada cola lo que minimiza el PV "
                       "(Murphy 2023). La revisión de ISO 10110-5 en preparación prevé admitir "
                       "PVr o PV99.8% en lugar del PV máx−mín salvo indicación en el plano."),
        'pvr_scope': ("* PVr fuera del alcance de ISO 10110-5:2015 (sólo apertura circular): "
                      "calculado con Zernike sobre el círculo que circunscribe la apertura."),
        'fast': ("Superficie rápida (semiapertura/R ≈ {s:.2f}): la esfera best-fit paraxial "
                 "puede desviar A del orden de un {p:.0f} %."),
        'iso_sag': ("Con «cambio de sag» los valores corresponden a la desviación de sagita ΔZ "
                    "(ISO 10110-5, 3.2); la norma define la desviación de forma normal a la "
                    "superficie (3.1). En espejos planos coinciden."),
        'n_pvr': ("El PV total ({pv:.3g} nm) es {r:.1f} veces el PV robusto PVr ({pvr:.3g} nm): "
                  "el PV está dominado por picos locales."),
        'n_iso_fail': "La superficie NO cumple la tolerancia ISO 10110-5 en: {k}.",
        'n_iso_ok': "La superficie cumple todas las tolerancias ISO 10110-5 comprobadas.",
        'file': "Fichero de datos",
        'case': "Caso de carga",
        'date': "Fecha del informe",
        'nodes': "Nodos válidos / leídos",
        'frame': "Sistema de referencia",
        'frame_global': "Global (Z del fichero = eje óptico)",
        'frame_bestfit': "Plano best-fit (normal inclinada {t:.3f}° respecto a Z)",
        'rigid': "Modos de sólido rígido eliminados",
        'none': "ninguno",
        'aperture': "Apertura",
        'ap_circ': "Circular · centro ({cx:.3f}, {cy:.3f}) · radio {r:.4g} {u}",
        'ap_ann': ("Anular · centro ({cx:.3f}, {cy:.3f}) · radios {ri:.4g} / {r:.4g} {u}"
                   " · obscuración ε = {e:.3f}"),
        'ap_hex': ("Hexagonal · centro ({cx:.3f}, {cy:.3f}) · circunradio {r:.4g} {u}"
                   " · vértices en {o}"),
        'ap_rect': "Rectangular · centro ({cx:.3f}, {cy:.3f}) · {w:.4g} × {h:.4g} {u}",
        'quantity': "Error de superficie",
        'q_sag': "Cambio de sag Δz = (u·n̂)/n_z (incluye el movimiento en el plano)",
        'q_normal': "Desplazamiento normal u·n̂",
        'fit': "Ajuste polinómico",
        'fit_desc': "{fam} · {n} términos · normalización {norm} · {method}",
        'no_fit': "No se ha realizado ajuste polinómico",
        'wls': "mínimos cuadrados ponderados por área (WLS)",
        'ls': "mínimos cuadrados ordinarios",
        'norm_pico': "pico", 'norm_rms': "RMS",
        'rbm_text': ("Se eliminó el movimiento de sólido rígido ({modes}) por mínimos "
                     "cuadrados ponderados por área sobre toda la superficie. El mapa "
                     "muestra el {qty} residual, es decir, la deformación de la "
                     "superficie sin traslaciones ni giros globales."),
        'rbm_none_text': ("No se eliminó ningún modo de sólido rígido: el mapa muestra el "
                          "{qty} tal como sale del modelo FEM."),
        'qty_sag': "cambio de sag Δz", 'qty_normal': "desplazamiento normal u·n̂",
        'rbm_values': "Movimiento rígido eliminado: {vals}.",
        'fit_text': ("Sobre la apertura {ap} se ajustaron y restaron los términos de la "
                     "tabla siguiente ({method}). El residuo es el error de figura que "
                     "no se corrige con esos términos."),
        'fig_rbm': "Figura {k}. Error de superficie tras eliminar el sólido rígido.",
        'fig_res': "Figura {k}. Residuo tras eliminar {n} términos de {fam}.",
        'tab_stats': "Resultados",
        'q': "Magnitud", 'v': "Valor",
        'rms': "RMS (ponderado por área)",
        'max': "Máximo (pico)", 'min': "Mínimo (valle)",
        'pv': "PV (máximo − mínimo)",
        'at': "en ({x:.3g}, {y:.3g})",
        'red': "Reducción de RMS respecto a la sección 2",
        'term': "Término", 'sc': "Nº", 'desc': "Descripción",
        'coef': "Coeficiente [{u}]",
        'cbar': "Error de superficie [{u}]", 'cbar_res': "Residuo [{u}]",
        'contrib': "Aporte RMS [{u}]", 'rms_after': "RMS residual [{u}]",
        'contrib_note': ("Aporte RMS: contribución de cada término al RMS, independiente del "
                         "número de términos ajustados (Gram–Schmidt sobre los nodos). RMS "
                         "residual: tras eliminar hasta ese término. Nº: índice 1-based "
                         "en la convención de la base."),
        'leg_max': "Máx.", 'leg_min': "Mín.", 'leg_ap': "Apertura",
        'n_reduction': "Eliminar los {n} términos reduce el RMS un {p:.1f} % "
                       "(de {a} a {b}).",
        'n_dom': "El término dominante entre los eliminados es {name} "
                 "(coeficiente {c}).",
        'n_spikes': ("El PV es {r:.0f} veces el RMS en la {where}: indica picos muy "
                     "localizados (por ejemplo, cargas puntuales o singularidades del "
                     "modelo), con el máximo en ({x:.3g}, {y:.3g}). El PV no es "
                     "representativo de la superficie global."),
        'where_rbm': "superficie completa", 'where_res': "superficie residual",
        'n_cover': ("La malla cubre solo un {p:.0f} % de la apertura teórica: puede "
                    "tratarse de un modelo parcial; revise el centro y el tamaño de la "
                    "apertura."),
        'n_excl': "{n} nodos se excluyeron por estar repetidos o fuera de la triangulación.",
        'n_def': ("Definiciones: RMS ponderado por área nodal; PV = máximo − mínimo sobre "
                  "los nodos. Los valores de RMS y PV del residuo no dependen de la "
                  "normalización de los polinomios."),
        'unit_note': "Valores en {u}, unidades del fichero de entrada.",
        'fam': {'fringe': "Zernike Fringe", 'noll': "Zernike Noll",
                'annular': "Zernike anulares (Mahajan)",
                'hexagonal': "polinomios hexagonales (Mahajan–Dai)",
                'rectangular': "polinomios rectangulares (Mahajan–Dai)"},
        'ap_name': {'circular': 'circular', 'annular': 'anular',
                    'hexagonal': 'hexagonal', 'rectangular': 'rectangular'},
    },
    'en': {
        'title': "FEM mirror surface analysis report",
        'header': "FEM analysis – mirror surface",
        'footer': "Generated with fem-mirror-analysis",
        'page': "Page",
        's_setup': "1. Analysis data",
        's_rbm': "2. Surface after rigid body motion removal",
        's_fit': "3. Residual after removing {n} {fam} terms",
        's_notes': "{k}. Remarks",
        's_opt': "{k}. Optical quality",
        's_iso': "{k}. ISO 10110-5 surface form tolerances",
        'combo': "Load case combination",
        'opt_text': ("Wavefront error in reflection WFE = 2·cos θ·(surface error), with "
                     "λ = {lam:g} nm and incidence θ = {inc:g}°, piston and tilt removed. Strehl "
                     "from the Maréchal approximation S ≈ exp[−(2πσ_W/λ)²] and numerically from "
                     "the PSF computed by FFT of the pupil."),
        'opt_surf': "Surface without rigid body motion", 'opt_res': "Fit residual",
        'opt_h': ["", "σ_W [nm]", "σ_W [λ]", "Strehl (Maréchal)"],
        'opt_fft': "Numerical Strehl (FFT PSF, {on})",
        'on_res': "residual", 'on_surf': "surface",
        'opt_note_lo': "Strehl < 0.3: the Maréchal approximation is not reliable; use the numerical value.",
        'opt_note_dl': ("Surface without rigid body motion has Strehl ≥ 0.8 (σ_W ≤ λ/14): "
                        "diffraction-limited at this wavelength (Maréchal criterion)."),
        'opt_note_ndl': ("Surface without rigid body motion has Strehl < 0.8: not "
                         "diffraction-limited at this λ unless the fitted terms are corrected."),
        'iso_text': ("Surface form deviation decomposed according to ISO 10110-5:2015 (Table B.2) "
                     "over the aperture, after tilt removal. Surface values in nm and in fringe "
                     "spacings (1 fringe = λ/2 = {fr:g} nm)."),
        'iso_h': ["Quantity", "Function", "nm", "fringes"],
        'iso_rows': [("A · power deviation", "PV(f_WS)"), ("B · irregularity", "PV(f_WI)"),
                     ("C · rot. invariant irregularity", "PV(f_WRI)"),
                     ("Total PV", "PV(f_WD)"), ("RMSt · total", "rms(f_WD)"),
                     ("RMSi · irregularity", "rms(f_WI)"),
                     ("RMSa · asymmetric irregularity", "rms(f_WRV)"),
                     ("PVr · robust PV", "PVr(f_WD)"),
                     ("PV99.8% · clipped PV", "PV%(f_WD)"),
                     ("PV99.8% · clipped irregularity", "PV%(f_WI)")],
        'iso_slope': "Slope of f_WI (FEM mesh sampling): RMS {r:.3g} µrad · max {m:.3g} µrad.",
        'iso_notation': "Measured values in ISO notation: {s}",
        'iso_check': "Check against tolerance {s}",
        'iso_ch_h': ["Quantity", "Measured [nm]", "Tolerance [nm]", "Result"],
        'ok': "Pass", 'nok': "Fail",
        'iso_method': ("Method: functions of ISO 14999-4:2007 (3.2–3.3) with RMS weighted by "
                       "nodal area. Tilt = best-fit plane (includes piston); best-fit sphere "
                       "approximated paraxially by the ρ² term; rotationally invariant part (whose "
                       "rms optimum is the azimuthal mean) approximated by even polynomials in ρ "
                       "up to ρ^{nr}; PVr after Evans (Opt. Eng. 48(4), 2009), PV of a 37-term "
                       "Fringe Zernike fit + 3·RMS of the residual, bounded between 6·RMS of the "
                       "residual and PV (Murphy, Proc. SPIE 12778, 2023); PV99.8% = PV of 99.8 % "
                       "of the area, clipping from each tail so as to minimise PV (Murphy 2023). "
                       "The ISO 10110-5 revision in preparation plans to allow PVr or PV99.8% "
                       "instead of max−min PV unless the drawing states otherwise."),
        'pvr_scope': ("* PVr outside the scope of ISO 10110-5:2015 (circular apertures only): "
                      "computed with Zernike over the circle circumscribing the aperture."),
        'fast': ("Fast surface (semi-aperture/R ≈ {s:.2f}): the paraxial best-fit sphere may "
                 "bias A by about {p:.0f} %."),
        'iso_sag': ("With “sag change” the values are sagitta deviation ΔZ (ISO 10110-5, 3.2); "
                    "the standard defines the surface form deviation normal to the surface (3.1). "
                    "Both coincide for flat mirrors."),
        'n_pvr': ("Total PV ({pv:.3g} nm) is {r:.1f} times the robust PVr ({pvr:.3g} nm): PV is "
                  "dominated by local peaks."),
        'n_iso_fail': "The surface does NOT meet the ISO 10110-5 tolerance on: {k}.",
        'n_iso_ok': "The surface meets all checked ISO 10110-5 tolerances.",
        'file': "Data file",
        'case': "Load case",
        'date': "Report date",
        'nodes': "Valid / read nodes",
        'frame': "Reference frame",
        'frame_global': "Global (file Z = optical axis)",
        'frame_bestfit': "Best-fit plane (normal tilted {t:.3f}° from Z)",
        'rigid': "Rigid body modes removed",
        'none': "none",
        'aperture': "Aperture",
        'ap_circ': "Circular · centre ({cx:.3f}, {cy:.3f}) · radius {r:.4g} {u}",
        'ap_ann': ("Annular · centre ({cx:.3f}, {cy:.3f}) · radii {ri:.4g} / {r:.4g} {u}"
                   " · obscuration ε = {e:.3f}"),
        'ap_hex': ("Hexagonal · centre ({cx:.3f}, {cy:.3f}) · circumradius {r:.4g} {u}"
                   " · vertices on {o}"),
        'ap_rect': "Rectangular · centre ({cx:.3f}, {cy:.3f}) · {w:.4g} × {h:.4g} {u}",
        'quantity': "Surface error",
        'q_sag': "Sag change Δz = (u·n̂)/n_z (includes in-plane motion)",
        'q_normal': "Normal displacement u·n̂",
        'fit': "Polynomial fit",
        'fit_desc': "{fam} · {n} terms · {norm} normalisation · {method}",
        'no_fit': "No polynomial fit performed",
        'wls': "area-weighted least squares (WLS)",
        'ls': "ordinary least squares",
        'norm_pico': "peak", 'norm_rms': "RMS",
        'rbm_text': ("Rigid body motion ({modes}) was removed by area-weighted least "
                     "squares over the whole surface. The map shows the residual {qty}, "
                     "i.e. the surface deformation without global translations or "
                     "rotations."),
        'rbm_none_text': ("No rigid body mode was removed: the map shows the {qty} as "
                          "given by the FEM model."),
        'qty_sag': "sag change Δz", 'qty_normal': "normal displacement u·n̂",
        'rbm_values': "Rigid body motion removed: {vals}.",
        'fit_text': ("Over the {ap} aperture, the terms in the table below were fitted "
                     "and subtracted ({method}). The residual is the figure error that "
                     "those terms cannot correct."),
        'fig_rbm': "Figure {k}. Surface error after rigid body motion removal.",
        'fig_res': "Figure {k}. Residual after removing {n} {fam} terms.",
        'tab_stats': "Results",
        'q': "Quantity", 'v': "Value",
        'rms': "RMS (area-weighted)",
        'max': "Maximum (peak)", 'min': "Minimum (valley)",
        'pv': "PV (maximum − minimum)",
        'at': "at ({x:.3g}, {y:.3g})",
        'red': "RMS reduction with respect to section 2",
        'term': "Term", 'sc': "No.", 'desc': "Description",
        'coef': "Coefficient [{u}]",
        'cbar': "Surface error [{u}]", 'cbar_res': "Residual [{u}]",
        'contrib': "RMS contribution [{u}]", 'rms_after': "Residual RMS [{u}]",
        'contrib_note': ("RMS contribution: contribution of each term to the RMS, independent "
                         "of the number of fitted terms (Gram–Schmidt over the nodes). "
                         "Residual RMS: after removing up to that term. No.: 1-based index "
                         "in the basis convention."),
        'leg_max': "Max.", 'leg_min': "Min.", 'leg_ap': "Aperture",
        'n_reduction': "Removing the {n} terms reduces the RMS by {p:.1f} % "
                       "(from {a} to {b}).",
        'n_dom': "The dominant removed term is {name} (coefficient {c}).",
        'n_spikes': ("PV is {r:.0f} times the RMS on the {where}: this indicates very "
                     "localised peaks (e.g. point loads or model singularities), with "
                     "the maximum at ({x:.3g}, {y:.3g}). PV is not representative of "
                     "the overall surface."),
        'where_rbm': "full surface", 'where_res': "residual surface",
        'n_cover': ("The mesh covers only {p:.0f} % of the nominal aperture: it may be a "
                    "partial model; check the aperture centre and size."),
        'n_excl': "{n} nodes were excluded (duplicated or outside the triangulation).",
        'n_def': ("Definitions: RMS weighted by nodal area; PV = maximum − minimum over "
                  "the nodes. Residual RMS and PV do not depend on the polynomial "
                  "normalisation."),
        'unit_note': "Values in {u}, the units of the input file.",
        'fam': {'fringe': "Fringe Zernike", 'noll': "Noll Zernike",
                'annular': "annular Zernike (Mahajan)",
                'hexagonal': "hexagonal polynomials (Mahajan–Dai)",
                'rectangular': "rectangular polynomials (Mahajan–Dai)"},
        'ap_name': {'circular': 'circular', 'annular': 'annular',
                    'hexagonal': 'hexagonal', 'rectangular': 'rectangular'},
    },
}

# Nombres en inglés de los términos (los españoles vienen de fem_mirror_analysis)
_ZERNIKE_EN = {
    (0, 0): "Piston", (1, 1): "Tilt X (cos θ)", (1, -1): "Tilt Y (sin θ)",
    (2, 0): "Defocus", (2, 2): "Astigmatism 0°/90° (cos 2θ)",
    (2, -2): "Astigmatism ±45° (sin 2θ)", (3, 1): "Coma X", (3, -1): "Coma Y",
    (3, 3): "Trefoil X (cos 3θ)", (3, -3): "Trefoil Y (sin 3θ)",
    (4, 0): "Primary spherical", (4, 2): "Secondary astigmatism 0°/90°",
    (4, -2): "Secondary astigmatism ±45°", (4, 4): "Tetrafoil (cos 4θ)",
    (4, -4): "Tetrafoil (sin 4θ)", (5, 1): "Secondary coma X",
    (5, -1): "Secondary coma Y", (5, 3): "Secondary trefoil X",
    (5, -3): "Secondary trefoil Y", (6, 0): "Secondary spherical",
    (8, 0): "Tertiary spherical",
}

_NM_FACTOR = {'m': 1e9, 'mm': 1e6, 'µm': 1e3, 'um': 1e3, 'nm': 1.0}


def _term_name(basis, t, lang):
    """Nombre del término (en las bases ortonormales, el del Zernike que lo genera)."""
    n, m = fma.term_nm(basis, t)
    if lang == 'en':
        return _ZERNIKE_EN.get((n, m), f"Z(n={n}, m={m:+d})")
    return fma.zernike_name(n, m)


# ═══════════════════════════════════════════════════════════════════════════════
#  DATOS DEL INFORME
# ═══════════════════════════════════════════════════════════════════════════════

def surface_stats(values, an, xy):
    """RMS ponderado, máximo, mínimo (con posición) y PV, ignorando NaN."""
    v = np.asarray(values, dtype=float)
    ok = np.isfinite(v)
    if not ok.any():
        return None
    idx = np.where(ok)[0]
    imax = idx[np.argmax(v[ok])]
    imin = idx[np.argmin(v[ok])]
    return {
        'rms': fma.rms_area_weighted(v, an),
        'max': float(v[imax]), 'xy_max': (float(xy[imax, 0]), float(xy[imax, 1])),
        'min': float(v[imin]), 'xy_min': (float(xy[imin, 0]), float(xy[imin, 1])),
        'pv': float(v[imax] - v[imin]),
    }


def build_report_data(params, coor, un, an, pos, rigid, mask, geom,
                      basis=None, terms=None, coeffs=None, residual=None,
                      wls=True, norm='pico', ndup=0, tilt_deg=0.0,
                      contrib=None, quantity='normal', optics=None):
    """
    Reúne todo lo necesario para el informe (se llama tras el análisis).

    coor, un, an : arrays de toda la superficie (sistema usado en el análisis)
    mask         : nodos de la apertura ; residual : residuo en esos nodos
    basis, terms, coeffs : None si no hubo ajuste
    """
    xy = coor[:, :2]
    valid = np.isfinite(un)
    data = {
        'fname': params.get('fname'),
        'case': params.get('case'),
        'n_read': int(len(coor)),
        'n_valid': int(valid.sum()),
        'n_excluded': int((~valid).sum()),
        'n_dup': int(ndup),
        'frame': params.get('frame', 'global'),
        'quantity': quantity,
        'tilt_deg': float(tilt_deg),
        'rigid': tuple(int(v) for v in rigid),
        'pos': [float(p) for p in pos],
        'x': xy[:, 0].copy(), 'y': xy[:, 1].copy(),
        'un': np.asarray(un, dtype=float).copy(),
        'stats_total': surface_stats(un, an, xy),
        'geom': dict(geom),
        'aperture': geom['type'],
        'n_ap': int(mask.sum()),
        'coverage': fma.mesh_coverage(xy, an, geom),
        'fit': None,
        'combo': params.get('combo_desc'),
        'optics': None,
    }
    if optics:
        data['optics'] = {k: optics[k] for k in ('iso', 'wfe_surface', 'wfe_residual', 'check',
                                                 'units', 'wavelength', 'incidence', 'spec')}
        data['optics']['tol_warnings'] = (optics['tol'] or {}).get('warnings', [])
        psf = optics.get('psf') or {}
        data['optics']['strehl_fft'] = psf.get('strehl')
        data['optics']['psf_on'] = psf.get('on')
    if basis is not None and terms is not None and coeffs is not None and residual is not None:
        res_full = np.full(len(coor), np.nan)
        res_full[mask] = residual
        an_ap = np.where(mask, an, 0.0)
        st_res = surface_stats(res_full, an_ap, xy)
        un_ap = np.where(mask, un, np.nan)
        rms_before = fma.rms_area_weighted(un_ap, an_ap)
        data['fit'] = {
            'basis': basis,
            'terms': [int(t) for t in terms],
            'coeffs': [float(c) for c in coeffs],
            'norm': norm,
            'wls': bool(wls),
            'residual': res_full,
            'stats': st_res,
            'rms_before': rms_before,
            'contrib': None if contrib is None else
                       [np.asarray(contrib[0]).tolist(), np.asarray(contrib[1]).tolist(),
                        np.asarray(contrib[2]).tolist()],
        }
    return data


# ═══════════════════════════════════════════════════════════════════════════════
#  FIGURAS
# ═══════════════════════════════════════════════════════════════════════════════

def _map_png(x, y, z, geom, cbar_label, units, T):
    """Mapa de color del campo z (NaN fuera) con contorno de apertura y máx/mín."""
    from matplotlib.figure import Figure

    ok = np.isfinite(z)
    fig = Figure(figsize=(6.4, 4.9), dpi=200)
    ax = fig.add_subplot(111)
    tri = fma.plot_triangulation(x[ok], y[ok])
    cf = ax.tricontourf(tri, z[ok], levels=64, cmap=fma.surface_cmap())
    cb = fig.colorbar(cf, ax=ax, shrink=0.9)
    cb.set_label(cbar_label)
    cb.formatter.set_powerlimits((-2, 3))
    cb.update_ticks()

    # Contorno de la apertura
    if geom:
        for xs, ys in fma.aperture_outline(geom):
            ax.plot(xs, ys, color='black', lw=1.0, ls='--', zorder=5)

    # Máximo y mínimo
    zi = np.where(ok)[0]
    imax, imin = zi[np.argmax(z[ok])], zi[np.argmin(z[ok])]
    ax.plot(x[imax], y[imax], marker='^', ms=8, mfc='white', mec='black',
            ls='none', zorder=6, label=T['leg_max'])
    ax.plot(x[imin], y[imin], marker='v', ms=8, mfc='black', mec='white',
            ls='none', zorder=6, label=T['leg_min'])
    ax.legend(loc='upper right', fontsize=7, framealpha=0.8)

    ax.set_aspect('equal')
    ax.set_xlabel(f'X [{units}]')
    ax.set_ylabel(f'Y [{units}]')
    ax.tick_params(labelsize=8)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format='png', bbox_inches='tight')
    buf.seek(0)
    return buf


# ═══════════════════════════════════════════════════════════════════════════════
#  UTILIDADES DOCX
# ═══════════════════════════════════════════════════════════════════════════════

def _fmt(v, units):
    """Valor en notación científica y, si se conocen las unidades, en nm."""
    if v is None or not np.isfinite(v):
        return '—'
    s = f"{v:.3e} {units}"
    f = _NM_FACTOR.get(units)
    if f is not None and units != 'nm':
        s += f"  ({v * f:.4g} nm)"
    return s


def _shade(cell, hex_fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def _table(doc, rows, header=None, widths=None, first_col_bold=True):
    from docx.shared import Pt, Cm
    ncols = len(rows[0]) if rows else len(header)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = 'Table Grid'
    if header:
        cells = t.add_row().cells
        for c, h in zip(cells, header):
            c.text = ''
            r = c.paragraphs[0].add_run(h)
            r.bold = True
            _shade(c, 'D9E2F3')
    for row in rows:
        cells = t.add_row().cells
        for k, (c, val) in enumerate(zip(cells, row)):
            c.text = ''
            r = c.paragraphs[0].add_run(str(val))
            if k == 0 and first_col_bold and not header:
                r.bold = True
                _shade(c, 'F2F2F2')
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def _add_page_field(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = '1'
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def _caption(doc, text):
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(8)


def _stats_rows(st, T, units, extra=None):
    rows = [
        [T['rms'], _fmt(st['rms'], units)],
        [T['max'], f"{_fmt(st['max'], units)}  {T['at'].format(x=st['xy_max'][0], y=st['xy_max'][1])}"],
        [T['min'], f"{_fmt(st['min'], units)}  {T['at'].format(x=st['xy_min'][0], y=st['xy_min'][1])}"],
        [T['pv'], _fmt(st['pv'], units)],
    ]
    if extra:
        rows.extend(extra)
    return rows


# ═══════════════════════════════════════════════════════════════════════════════
#  INFORME
# ═══════════════════════════════════════════════════════════════════════════════

def generate_report(data, path, lang='es', units='mm', logo_png=None):
    """
    Genera el informe Word en `path`.

    data     : dict de build_report_data
    lang     : 'es' | 'en'
    units    : unidades de longitud del fichero de entrada (p. ej. 'mm')
    logo_png : bytes PNG opcional para la cabecera
    """
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    T = TXT['en' if lang == 'en' else 'es']
    units = (units or '').strip() or '-'
    fit = data.get('fit')

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ('left_margin', 'right_margin'):
        setattr(sec, side, Cm(2.0))
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)

    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    for name, size in (('Title', 18), ('Heading 1', 13)):
        s = doc.styles[name]
        s.font.name = 'Calibri'
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor(0x1F, 0x3A, 0x68)

    # Cabecera y pie
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if logo_png:
        try:
            hp.add_run().add_picture(io.BytesIO(logo_png), height=Cm(1.0))
            hp.add_run('   ')
        except Exception:
            pass
    hr = hp.add_run(T['header'])
    hr.font.size = Pt(8)
    hr.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run(f"{T['footer']} · {T['page']} ")
    fr.font.size = Pt(8)
    _add_page_field(fp)

    # Título
    doc.add_paragraph(T['title'], style='Title')

    # ── 1. Datos del análisis ────────────────────────────────────────────────
    doc.add_heading(T['s_setup'], level=1)
    g = data['geom']
    cx, cy = [0.0 if abs(v) < 5e-4 else v for v in g['center']]   # evita "-0.000"
    if g['type'] == 'circular':
        ap = T['ap_circ'].format(cx=cx, cy=cy, r=g['radius'], u=units)
    elif g['type'] == 'annular':
        ap = T['ap_ann'].format(cx=cx, cy=cy, r=g['radius'], ri=g['inner_radius'],
                                e=g['eps'], u=units)
    elif g['type'] == 'hexagonal':
        ap = T['ap_hex'].format(cx=cx, cy=cy, r=g['radius'], u=units,
                                o='±X' if abs(g.get('rotation', 0)) < 1 else '±Y')
    else:
        ap = T['ap_rect'].format(cx=cx, cy=cy, w=g['width'], h=g['height'], u=units)
    modes = [m for m, f in zip(['Tx', 'Ty', 'Tz', 'Rx', 'Ry', 'Rz'], data['rigid']) if f]
    if fit:
        fam = T['fam'][fit['basis']]
        fit_txt = T['fit_desc'].format(
            fam=fam, n=len(fit['terms']), norm=T['norm_' + fit['norm']],
            method=T['wls'] if fit['wls'] else T['ls'])
    else:
        fam, fit_txt = '', T['no_fit']
    rows = [[T['file'], os.path.basename(data['fname'] or '—')]]
    if data.get('case') is not None:
        rows.append([T['case'], f"{data['case']:g}"])
    if data.get('combo'):
        rows.append([T['combo'], ' · '.join(data['combo'])])
    rows += [
        [T['date'], _dt.datetime.now().strftime('%Y-%m-%d %H:%M')],
        [T['nodes'], f"{data['n_valid']:,} / {data['n_read']:,}"],
        [T['frame'], T['frame_global'] if data['frame'] != 'bestfit'
                     else T['frame_bestfit'].format(t=data['tilt_deg'])],
        [T['rigid'], ', '.join(modes) if modes else T['none']],
        [T['quantity'], T['q_' + data.get('quantity', 'normal')]],
        [T['aperture'], ap],
        [T['fit'], fit_txt],
    ]
    _table(doc, rows, widths=[5.0, 12.0])
    p = doc.add_paragraph(T['unit_note'].format(u=units))
    p.runs[0].font.size = Pt(9)
    p.runs[0].italic = True

    # ── 2. Superficie sin sólido rígido ──────────────────────────────────────
    k_fig = 1
    doc.add_heading(T['s_rbm'], level=1)
    if modes:
        doc.add_paragraph(T['rbm_text'].format(modes=', '.join(modes),
                                               qty=T['qty_' + data.get('quantity', 'normal')]))
        vals = []
        for i, m in enumerate(['Tx', 'Ty', 'Tz', 'Rx', 'Ry', 'Rz']):
            if data['rigid'][i]:
                u = units if i < 3 else 'rad'
                vals.append(f"{m} = {data['pos'][i]:+.3e} {u}")
        pr = doc.add_paragraph(T['rbm_values'].format(vals='; '.join(vals)))
        pr.runs[0].font.size = Pt(9)
    else:
        doc.add_paragraph(T['rbm_none_text'].format(qty=T['qty_' + data.get('quantity', 'normal')]))
    img = _map_png(data['x'], data['y'], data['un'], g,
                   T['cbar'].format(u=units), units, T)
    doc.add_picture(img, width=Cm(12.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _caption(doc, T['fig_rbm'].format(k=k_fig))
    k_fig += 1
    _table(doc, _stats_rows(data['stats_total'], T, units), widths=[5.0, 12.0])

    # ── 3. Residuo del ajuste ────────────────────────────────────────────────
    if fit:
        n = len(fit['terms'])
        doc.add_heading(T['s_fit'].format(n=n, fam=fam), level=1)
        doc.add_paragraph(T['fit_text'].format(
            ap=T['ap_name'].get(data['aperture'], data['aperture']),
            method=T['wls'] if fit['wls'] else T['ls']))
        trows = []
        con = fit.get('contrib')
        for k, (t, c) in enumerate(zip(fit['terms'], fit['coeffs'])):
            lab, sc, _, _ = fma.term_info(fit['basis'], t)
            row = [lab, str(sc), _term_name(fit['basis'], t, lang), f"{c:+.3e}"]
            if con:
                row += [f"{con[0][k]:.3e}", f"{con[1][k]:.3e}"]
            trows.append(row)
        header = [T['term'], T['sc'], T['desc'], T['coef'].format(u=units)]
        widths = [1.5, 1.1, 6.6, 2.7]
        if con:
            header += [T['contrib'].format(u=units), T['rms_after'].format(u=units)]
            widths += [2.6, 2.7]
        _table(doc, trows, header=header, widths=widths)
        if con:
            pn = doc.add_paragraph(T['contrib_note'])
            pn.runs[0].font.size = Pt(8.5)
            pn.runs[0].italic = True
        img = _map_png(data['x'], data['y'], fit['residual'], g,
                       T['cbar_res'].format(u=units), units, T)
        doc.add_picture(img, width=Cm(12.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _caption(doc, T['fig_res'].format(k=k_fig, n=n, fam=fam))
        k_fig += 1
        rb, ra = fit['rms_before'], fit['stats']['rms']
        red = (1 - ra / rb) * 100 if rb and rb > 0 else float('nan')
        _table(doc, _stats_rows(fit['stats'], T, units,
                                extra=[[T['red'], f"{red:.1f} %"]]),
               widths=[5.0, 12.0])

    # ── Calidad óptica e ISO 10110-5 ─────────────────────────────────────────
    k_sec = 4 if fit else 3
    opt = data.get('optics')
    if opt:
        lam = opt['wavelength']
        doc.add_heading(T['s_opt'].format(k=k_sec), level=1)
        k_sec += 1
        doc.add_paragraph(T['opt_text'].format(lam=lam, inc=opt['incidence']))
        orows = []
        for key, lab in (('wfe_surface', 'opt_surf'), ('wfe_residual', 'opt_res')):
            oq = opt.get(key)
            if oq:
                orows.append([T[lab], f"{oq['wfe_rms_nm']:.3f}",
                              f"{oq['wfe_rms_waves']:.4f}" + (f"  (λ/{1/oq['wfe_rms_waves']:.0f})"
                                                              if oq['wfe_rms_waves'] > 0 else ""),
                              f"{oq['strehl']:.4f}"])
        _table(doc, orows, header=T['opt_h'], widths=[6.0, 3.0, 4.0, 4.0])
        if opt.get('strehl_fft') is not None:
            on = T['on_res'] if opt.get('psf_on') == 'residual' else T['on_surf']
            doc.add_paragraph(f"{T['opt_fft'].format(on=on)}: {opt['strehl_fft']:.4f}")

        iso = opt['iso']
        fr = lam / 2.0
        doc.add_heading(T['s_iso'].format(k=k_sec), level=1)
        k_sec += 1
        doc.add_paragraph(T['iso_text'].format(fr=fr))
        keys = ['A', 'B', 'C', 'PV', 'RMSt', 'RMSi', 'RMSa', 'PVr', 'PV%', 'PV%_I']
        irows = []
        for key, (lab, fn) in zip(keys, T['iso_rows']):
            v = iso.get(key)
            if v is None:
                continue
            if key == 'PVr' and not iso.get('PVr_in_scope', True):
                lab += ' *'
            irows.append([lab, fn, f"{v:.3f}", f"{v / fr:.4f}"])
        _table(doc, irows, header=T['iso_h'], widths=[6.5, 3.0, 3.5, 3.5])
        extra_notes = []
        if iso.get('PVr') is not None and not iso.get('PVr_in_scope', True):
            extra_notes.append(T['pvr_scope'])
        if iso.get('sphere_s') and iso['sphere_s'] ** 2 / 4 > 0.01:
            extra_notes.append(T['fast'].format(s=iso['sphere_s'], p=100 * iso['sphere_s'] ** 2 / 4))
        for en_ in extra_notes:
            pe = doc.add_paragraph(en_)
            pe.runs[0].font.size = Pt(9)
        if iso.get('slope_rms_urad') is not None:
            doc.add_paragraph(T['iso_slope'].format(r=iso['slope_rms_urad'], m=iso['slope_max_urad']))
        comma = lang != 'en'
        doc.add_paragraph(T['iso_notation'].format(s=fma.format_iso10110(iso, comma=comma)))
        if opt.get('spec'):
            doc.add_paragraph(T['iso_check'].format(s=opt['spec']))
            for wmsg in opt.get('tol_warnings', []):
                pw = doc.add_paragraph(f"⚠ {wmsg}")
                pw.runs[0].font.size = Pt(9)
            if opt['check']:
                _table(doc, [[k_, f"{m_:.3f}", f"{l_:.3f}", T['ok'] if ok_ else T['nok']]
                             for k_, m_, l_, ok_ in opt['check']],
                       header=T['iso_ch_h'], widths=[3.5, 4.0, 4.0, 4.0])
        for key in ('iso_method',) + (('iso_sag',) if data.get('quantity') == 'sag' else ()):
            pm = doc.add_paragraph(T[key].format(nr=2 * iso.get('n_radial', 12)))
            pm.runs[0].font.size = Pt(8.5)
            pm.runs[0].italic = True

    # ── Observaciones ────────────────────────────────────────────────────────
    notes = []
    if fit:
        rb, ra = fit['rms_before'], fit['stats']['rms']
        if rb and rb > 0:
            notes.append(T['n_reduction'].format(n=len(fit['terms']),
                                                 p=(1 - ra / rb) * 100,
                                                 a=f"{rb:.3e} {units}", b=f"{ra:.3e} {units}"))
        con = fit.get('contrib')
        size = con[0] if con else [abs(c) for c in fit['coeffs']]
        cand = [(sz, t, c) for sz, t, c in zip(size, fit['terms'], fit['coeffs'])
                if fma.term_nm(fit['basis'], t) != (0, 0)]
        if cand:
            _, t, c = max(cand)
            notes.append(T['n_dom'].format(name=_term_name(fit['basis'], t, lang),
                                           c=f"{c:+.3e} {units}"))
    # Picos localizados: una sola observación (la primera superficie que los muestre)
    for st_, where in ((data['stats_total'], 'where_rbm'),
                       (fit['stats'] if fit else None, 'where_res')):
        ref = data['stats_total']['rms'] if data.get('stats_total') else 0.0
        if st_ and st_['rms'] > 1e-6 * ref and st_['pv'] / st_['rms'] > 10:
            notes.append(T['n_spikes'].format(r=st_['pv'] / st_['rms'], where=T[where],
                                              x=st_['xy_max'][0], y=st_['xy_max'][1]))
            break
    if opt:
        strehls = [o['strehl'] for o in (opt.get('wfe_surface'), opt.get('wfe_residual')) if o]
        smin = min(strehls)
        if smin < 0.3:
            notes.append(T['opt_note_lo'])
        s_ref = opt['wfe_surface']['strehl']
        notes.append(T['opt_note_dl'] if s_ref >= 0.8 else T['opt_note_ndl'])
        iso = opt['iso']
        if iso.get('PVr') and iso['PVr'] > 0 and iso['PV'] / iso['PVr'] > 1.5:
            notes.append(T['n_pvr'].format(pv=iso['PV'], r=iso['PV'] / iso['PVr'], pvr=iso['PVr']))
        if opt.get('check'):
            bad = [k_ for k_, _, _, ok_ in opt['check'] if not ok_]
            notes.append(T['n_iso_fail'].format(k=', '.join(bad)) if bad else T['n_iso_ok'])
    if data.get('coverage') is not None and data['coverage'] < 0.82:
        notes.append(T['n_cover'].format(p=100 * data['coverage']))
    if data['n_excluded']:
        notes.append(T['n_excl'].format(n=data['n_excluded']))
    notes.append(T['n_def'])

    h = doc.add_heading(T['s_notes'].format(k=k_sec), level=1)
    h.paragraph_format.keep_with_next = True
    for n_ in notes:
        doc.add_paragraph(n_, style='List Bullet')

    doc.save(path)
    return path
